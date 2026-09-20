import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from PIL import Image, ImageTk
import cv2
import numpy as np
import time
import json
from datetime import datetime
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from sklearn.metrics import confusion_matrix, accuracy_score, precision_score, recall_score, f1_score

from utils.config import Config
from utils.logger import Logger
from core.camera_handler import CameraHandler
from core.face_recognizer import FaceRecognizer
from core.face_detector import FaceDetector
from core.trainer import FaceTrainer


class AnalysisGUI:
    """GUI untuk analisis pengenalan wajah"""

    def __init__(self, parent):
        self.parent = parent
        self.window = None
        self.camera = None
        self.recognizer = FaceRecognizer()
        self.detector = FaceDetector()
        self.trainer = FaceTrainer()
        self.logger = Logger("AnalysisGUI")

        self.is_running = False
        self.is_analyzing = False

        # Data analisis
        self.analysis_data = []
        self.analysis_counter = 1

        # Foto upload
        self.uploaded_photo = None
        self.uploaded_photo_path = None
        self.uploaded_photo_name = ""
        self.uploaded_face_roi = None

        # Preview mode
        self.preview_mode = 'camera'
        self.preview_label = None
        self.photo_preview_label = None

        # Table
        self.result_tree = None
        self.registered_users = []
        self.total_test_label = None
        self.status_label = None

        # UI Elements
        self.name_combo = None
        self.name_var = None
        self.type_var = None
        self.distance_var = None
        self.start_camera_btn = None
        self.stop_camera_btn = None
        self.analyze_btn = None
        self.camera_preview_btn = None
        self.photo_preview_btn = None
        self.mode_label = None
        self.photo_name_label = None

        # Optimasi preview
        self.preview_interval = 66
        self._preview_after_id = None
        self._last_frame_hash = None

        # Path untuk menyimpan data
        self.data_file = Config.DATA_DIR / "analysis_data.json"
        self.load_analysis_data()

    # ========================================================================
    # STATUS LABEL (TP/FP/TN/FN)
    # ========================================================================

    def _get_status_label(self, result):
        """
        Menentukan status label berdasarkan hasil pengenalan

        Parameters:
        - result: dict dengan keys 'type', 'expected_name', 'recognized_name', 'face_detected'

        Returns:
        - status_code: string (TP/FP/TN/FN)
        - status_text: string (deskripsi lengkap)
        """
        user_type = result['type']
        expected_name = result['expected_name']
        recognized_name = result['recognized_name']
        face_detected = result.get('face_detected', True)

        if user_type == "Terdaftar":
            if recognized_name not in ["Unknown", "No Face", "Error"] and face_detected:
                if recognized_name == expected_name:
                    return "TP", "Berhasil Dikenali"
                else:
                    return "FN", "Gagal (Ditolak)"
            else:
                return "FN", "Gagal (Ditolak)"
        else:
            if recognized_name not in ["Unknown", "No Face", "Error"] and face_detected:
                return "FP", "Gagal (Dikenali)"
            else:
                return "TN", "Berhasil Ditolak"

    # ========================================================================
    # GUI SETUP
    # ========================================================================

    def show(self):
        """Menampilkan GUI analisis"""
        self.window = tk.Toplevel(self.parent)
        self.window.title("Analisis Performa Pengenalan Wajah")
        self.window.geometry("1250x650")
        self.window.configure(bg='#f5f5f5')
        self.window.resizable(False, False)

        # Center window
        self.window.update_idletasks()
        width = 1250
        height = 650
        x = (self.window.winfo_screenwidth() // 2) - (width // 2)
        y = (self.window.winfo_screenheight() // 2) - (height // 2)
        self.window.geometry(f'{width}x{height}+{x}+{y}')

        self.window.protocol("WM_DELETE_WINDOW", self.on_closing)

        # Load registered users
        self.load_registered_users()
        self.recognizer.initialize()

        # Main container
        main_container = tk.Frame(self.window, bg='#f5f5f5')
        main_container.pack(fill='both', expand=True, padx=10, pady=10)

        # ========== TOP PANEL - KONTROL PENGUJIAN ==========
        top_panel = tk.Frame(main_container, bg='#ffffff', relief=tk.RIDGE, bd=1)
        top_panel.pack(fill='x', pady=(0, 10))

        # Title
        title_frame = tk.Frame(top_panel, bg='#FF9800', height=35)
        title_frame.pack(fill='x')
        title_frame.pack_propagate(False)

        tk.Label(title_frame, text="Kontrol Pengujian", font=('Arial', 12, 'bold'),
                 bg='#FF9800', fg='white').pack(pady=6)

        # Control content
        control_content = tk.Frame(top_panel, bg='#ffffff', padx=15, pady=10)
        control_content.pack(fill='x')

        # Row 1: Jarak, Tipe
        row1 = tk.Frame(control_content, bg='#ffffff')
        row1.pack(fill='x', pady=3)

        tk.Label(row1, text="Jarak (m):", font=('Arial', 10, 'bold'), bg='#ffffff', width=10, anchor='w').pack(
            side='left')
        self.distance_var = tk.StringVar(value="1.0")
        ttk.Combobox(row1, textvariable=self.distance_var, values=["0.5", "1.0", "1.5", "2.0", "2.5"],
                     state="readonly", width=10, font=('Arial', 10)).pack(side='left', padx=5)

        tk.Label(row1, text="Tipe:", font=('Arial', 10, 'bold'), bg='#ffffff', width=5, anchor='w').pack(side='left',
                                                                                                         padx=(30, 0))
        self.type_var = tk.StringVar(value="Terdaftar")
        type_frame = tk.Frame(row1, bg='#ffffff')
        type_frame.pack(side='left', padx=5)
        tk.Radiobutton(type_frame, text="Terdaftar", variable=self.type_var, value="Terdaftar",
                       bg='#ffffff', font=('Arial', 9), command=self.on_type_change).pack(side='left')
        tk.Radiobutton(type_frame, text="Tidak Terdaftar", variable=self.type_var, value="Tidak Terdaftar",
                       bg='#ffffff', font=('Arial', 9), command=self.on_type_change).pack(side='left', padx=(10, 0))

        # Row 2: Subjek, Mode Upload
        row2 = tk.Frame(control_content, bg='#ffffff')
        row2.pack(fill='x', pady=3)

        tk.Label(row2, text="Subjek:", font=('Arial', 10, 'bold'), bg='#ffffff', width=10, anchor='w').pack(side='left')
        self.name_var = tk.StringVar()
        self.name_combo = ttk.Combobox(row2, textvariable=self.name_var, values=self.registered_users,
                                       state="readonly", width=20, font=('Arial', 10))
        self.name_combo.pack(side='left', padx=5)
        if self.registered_users:
            self.name_var.set(self.registered_users[0])

        # Mode indicator
        self.mode_frame = tk.Frame(row2, bg='#e3f2fd', relief=tk.RIDGE, bd=1)
        self.mode_frame.pack(side='left', padx=(30, 5))
        self.mode_label = tk.Label(self.mode_frame, text="[Mode Kamera]", font=('Arial', 9, 'bold'),
                                   bg='#e3f2fd', fg='#1565C0')
        self.mode_label.pack(padx=10, pady=2)

        self.photo_name_label = tk.Label(row2, text="", font=('Arial', 9), bg='#ffffff', fg='#666666')
        self.photo_name_label.pack(side='left', padx=5)

        # Row 3: Tombol-tombol
        row3 = tk.Frame(control_content, bg='#ffffff')
        row3.pack(fill='x', pady=5)

        self.start_camera_btn = tk.Button(row3, text="▶ Mulai Kamera", font=('Arial', 9),
                                          bg='#4CAF50', fg='white', padx=12, pady=4,
                                          command=self.start_camera)
        self.start_camera_btn.pack(side='left', padx=2)

        self.stop_camera_btn = tk.Button(row3, text="■ Stop Kamera", font=('Arial', 9),
                                         bg='#f44336', fg='white', padx=12, pady=4,
                                         command=self.stop_camera, state='disabled')
        self.stop_camera_btn.pack(side='left', padx=2)

        tk.Button(row3, text="📷 Upload Foto", font=('Arial', 9),
                  bg='#2196F3', fg='white', padx=12, pady=4,
                  command=self.upload_photo).pack(side='left', padx=2)

        self.analyze_btn = tk.Button(row3, text="▶ Mulai Pengujian", font=('Arial', 9, 'bold'),
                                     bg='#FF9800', fg='white', padx=15, pady=4,
                                     command=self.start_analysis, state='disabled')
        self.analyze_btn.pack(side='left', padx=10)

        tk.Button(row3, text="🔄 Reset Data", font=('Arial', 9),
                  bg='#9E9E9E', fg='white', padx=12, pady=4,
                  command=self.delete_all_rows).pack(side='left', padx=2)

        tk.Button(row3, text="📊 Export Excel", font=('Arial', 9, 'bold'),
                  bg='#4CAF50', fg='white', padx=12, pady=4,
                  command=self.export_to_excel).pack(side='left', padx=2)

        # Row 4: Status dan Info
        row4 = tk.Frame(control_content, bg='#ffffff')
        row4.pack(fill='x', pady=3)

        self.status_label = tk.Label(row4, text="Status: Siap", font=('Arial', 9),
                                     bg='#ffffff', fg='#2196F3')
        self.status_label.pack(side='left')

        self.total_test_label = tk.Label(row4, text="Total Data: 0 pengujian", font=('Arial', 9),
                                         bg='#ffffff', fg='#666666')
        self.total_test_label.pack(side='right')

        # Info sistem
        info_text = f"Kamera: {Config.CAMERA_WIDTH}x{Config.CAMERA_HEIGHT} | Preview: 640x360 | Threshold: {Config.LBPH_THRESHOLD}"
        tk.Label(control_content, text=info_text, font=('Arial', 8),
                 bg='#ffffff', fg='#999999').pack(pady=3)

        # ========== BOTTOM PANEL - PREVIEW & HASIL ==========
        bottom_panel = tk.Frame(main_container, bg='#f5f5f5')
        bottom_panel.pack(fill='both', expand=True)

        # ===== LEFT - PREVIEW =====
        left_panel = tk.Frame(bottom_panel, bg='#ffffff', relief=tk.RIDGE, bd=1, width=680)
        left_panel.pack(side='left', fill='y', padx=(0, 5))
        left_panel.pack_propagate(False)

        preview_title = tk.Frame(left_panel, bg='#FF9800', height=35)
        preview_title.pack(fill='x')
        preview_title.pack_propagate(False)

        tk.Label(preview_title, text="Preview", font=('Arial', 12, 'bold'),
                 bg='#FF9800', fg='white').pack(pady=6)

        preview_container = tk.Frame(left_panel, bg='#000000')
        preview_container.pack(pady=10)

        self.preview_label = tk.Label(preview_container, bg='#000000', width=640, height=360)
        self.preview_label.pack()
        self.preview_label.pack_propagate(False)

        # Toggle buttons
        toggle_frame = tk.Frame(left_panel, bg='#ffffff', pady=5)
        toggle_frame.pack(fill='x')

        self.camera_preview_btn = tk.Button(toggle_frame, text="Kamera", font=('Arial', 9, 'bold'),
                                            bg='#FF9800', fg='white', padx=15, pady=3,
                                            command=self.show_camera_preview)
        self.camera_preview_btn.pack(side='left', padx=5)

        self.photo_preview_btn = tk.Button(toggle_frame, text="Foto", font=('Arial', 9),
                                           bg='#9E9E9E', fg='white', padx=15, pady=3,
                                           command=self.show_photo_preview)
        self.photo_preview_btn.pack(side='left', padx=5)

        tk.Label(toggle_frame, text="Preview: 640x360",
                 font=('Arial', 8), bg='#ffffff', fg='#999999').pack(side='right', padx=10)

        # ===== RIGHT - HASIL PENGUJIAN =====
        right_panel = tk.Frame(bottom_panel, bg='#ffffff', relief=tk.RIDGE, bd=1)
        right_panel.pack(side='right', fill='both', expand=True, padx=(5, 0))

        result_title = tk.Frame(right_panel, bg='#4CAF50', height=35)
        result_title.pack(fill='x')
        result_title.pack_propagate(False)

        tk.Label(result_title, text="Hasil Pengujian", font=('Arial', 12, 'bold'),
                 bg='#4CAF50', fg='white').pack(pady=6)

        # Treeview frame dengan scrollbar
        tree_frame = tk.Frame(right_panel, bg='#ffffff', padx=5, pady=5)
        tree_frame.pack(fill='both', expand=True)

        tree_scroll_y = tk.Scrollbar(tree_frame, orient='vertical')
        tree_scroll_y.pack(side='right', fill='y')

        tree_scroll_x = tk.Scrollbar(tree_frame, orient='horizontal')
        tree_scroll_x.pack(side='bottom', fill='x')

        # Kolom dengan status TP/FP/TN/FN
        columns = ('No', 'Jarak', 'Tipe', 'Subjek', 'Hasil', 'Conf', 'Waktu', 'Status', 'Kode')
        self.result_tree = ttk.Treeview(tree_frame, columns=columns, show='headings',
                                        yscrollcommand=tree_scroll_y.set,
                                        xscrollcommand=tree_scroll_x.set,
                                        height=12)

        tree_scroll_y.config(command=self.result_tree.yview)
        tree_scroll_x.config(command=self.result_tree.xview)

        # Lebar kolom
        col_widths = {
            'No': 45,
            'Jarak': 65,
            'Tipe': 100,
            'Subjek': 140,
            'Hasil': 140,
            'Conf': 80,
            'Waktu': 85,
            'Status': 160,
            'Kode': 55
        }

        for col in columns:
            self.result_tree.heading(col, text=col)
            self.result_tree.column(col, width=col_widths.get(col, 100), anchor='center', minwidth=50)

        self.result_tree.pack(fill='both', expand=True)

        # Load data setelah treeview siap
        self.load_data_to_treeview()
        self.total_test_label.config(text=f"Total Data: {len(self.analysis_data)} pengujian")

        # Action buttons
        action_frame = tk.Frame(right_panel, bg='#ffffff', pady=5)
        action_frame.pack(fill='x')

        tk.Button(action_frame, text="Hapus Baris", font=('Arial', 9),
                  bg='#ff9800', fg='white', padx=10, pady=3,
                  command=self.delete_selected_row).pack(side='left', padx=5)

        tk.Button(action_frame, text="Hapus Semua", font=('Arial', 9),
                  bg='#f44336', fg='white', padx=10, pady=3,
                  command=self.delete_all_rows).pack(side='left', padx=5)

        # Legenda status
        legend_frame = tk.Frame(right_panel, bg='#f5f5f5', pady=3)
        legend_frame.pack(fill='x')

        legend_text = "TP=Berhasil Dikenali | FP=Gagal (Dikenali) | TN=Berhasil Ditolak | FN=Gagal (Ditolak)"
        tk.Label(legend_frame, text=legend_text, font=('Arial', 7),
                 bg='#f5f5f5', fg='#666666').pack()

        self.update_button_states()
        self.start_preview_loop()

    # ========================================================================
    # PREVIEW LOOP
    # ========================================================================

    def start_preview_loop(self):
        """Memulai preview loop"""
        self._schedule_preview_update()

    def _schedule_preview_update(self):
        """Schedule preview update berikutnya"""
        if self.window:
            self._preview_after_id = self.window.after(self.preview_interval, self._preview_update)

    def _preview_update(self):
        """Update preview (internal method)"""
        if not self.is_running or not self.camera or not self.camera.is_available():
            self._schedule_preview_update()
            return

        if self.preview_mode == 'camera':
            frame = self.camera.get_frame()
            if frame is not None:
                frame_hash = hash(frame.tobytes()[:1000])
                if frame_hash == self._last_frame_hash:
                    self._schedule_preview_update()
                    return
                self._last_frame_hash = frame_hash

                face = self.detector.get_largest_face(frame)
                if face is not None:
                    self.detector.draw_face_box(frame, face, (0, 255, 0))

                display_frame = cv2.resize(frame, (640, 360))
                frame_rgb = cv2.cvtColor(display_frame, cv2.COLOR_BGR2RGB)
                img = Image.fromarray(frame_rgb)
                photo = ImageTk.PhotoImage(img)

                self.preview_label.config(image=photo, bg='#000000')
                self.preview_label.image = photo

        self._schedule_preview_update()

    # ========================================================================
    # DATA LOADING
    # ========================================================================

    def load_registered_users(self):
        """Memuat daftar user terdaftar"""
        self.trainer.load_model()
        self.registered_users = self.trainer.get_all_users()

    def load_data_to_treeview(self):
        """Memuat data dari memory ke treeview"""
        if self.result_tree is None:
            return

        self.result_tree.delete(*self.result_tree.get_children())

        # Setup tag colors
        self.result_tree.tag_configure('tp', background='#c8e6c9')
        self.result_tree.tag_configure('tn', background='#bbdefb')
        self.result_tree.tag_configure('fp', background='#ffcdd2')
        self.result_tree.tag_configure('fn', background='#fff9c4')

        for result in self.analysis_data:
            status_code, status_text = self._get_status_label(result)

            expected_name = result['expected_name']
            recognized_name = result['recognized_name']
            if len(expected_name) > 15:
                expected_name = expected_name[:15] + "..."
            if len(recognized_name) > 15:
                recognized_name = recognized_name[:15] + "..."

            self.result_tree.insert('', 'end', values=(
                result['no'],
                f"{result['distance']:.1f}m",
                result['type'][:15],
                expected_name,
                recognized_name,
                f"{result['confidence_percent']:.1f}%",
                f"{result['compute_time_ms']:.1f}ms",
                status_text,
                status_code
            ), tags=(status_code.lower(),))

    # ========================================================================
    # UI EVENT HANDLERS
    # ========================================================================

    def on_type_change(self):
        """Handler saat tipe berubah"""
        if self.type_var.get() == "Terdaftar":
            self.name_combo.config(values=self.registered_users, state="readonly")
            if self.registered_users:
                self.name_var.set(self.registered_users[0])
        else:
            self.name_combo.config(values=[], state="disabled")
            self.name_var.set("orang asing")

    def start_camera(self):
        """Memulai kamera"""
        self.camera = CameraHandler()
        if self.camera.start():
            self.is_running = True
            self._last_frame_hash = None
            self.start_camera_btn.config(state='disabled')
            self.stop_camera_btn.config(state='normal')
            self.analyze_btn.config(state='normal')
            self.status_label.config(text="Status: Kamera aktif", fg='#4CAF50')
            self.preview_mode = 'camera'
            self.update_button_states()
            self.mode_label.config(text="[Mode Kamera]")
            self.photo_name_label.config(text="")
        else:
            messagebox.showerror("Error", "Tidak dapat membuka kamera!")

    def stop_camera(self):
        """Menghentikan kamera"""
        self.is_running = False
        self._last_frame_hash = None
        if self.camera:
            self.camera.stop()
            self.camera = None
        self.start_camera_btn.config(state='normal')
        self.stop_camera_btn.config(state='disabled')
        self.analyze_btn.config(state='disabled')
        self.status_label.config(text="Status: Kamera tidak aktif", fg='#f44336')
        self.preview_label.config(image='', bg='#000000')

    def show_camera_preview(self):
        """Tampilkan preview kamera"""
        self.preview_mode = 'camera'
        self.update_button_states()
        self.mode_label.config(text="[Mode Kamera]")
        self.photo_name_label.config(text="")
        if not self.is_running:
            self.preview_label.config(image='', bg='#000000')

    def show_photo_preview(self):
        """Tampilkan preview foto"""
        self.preview_mode = 'photo'
        self.update_button_states()
        if self.uploaded_photo is not None:
            self.mode_label.config(text="[Mode Upload]")
            self.photo_name_label.config(text=self.uploaded_photo_name)
            display_img = self.uploaded_photo.copy()
            face = self.detector.get_largest_face(display_img)
            if face is not None:
                self.detector.draw_face_box(display_img, face, (255, 0, 0))

            display_img = cv2.resize(display_img, (640, 360))
            img_rgb = cv2.cvtColor(display_img, cv2.COLOR_BGR2RGB)
            photo = ImageTk.PhotoImage(Image.fromarray(img_rgb))

            self.preview_label.config(image=photo, bg='#000000')
            self.preview_label.image = photo
        else:
            self.preview_label.config(image='', bg='#e0e0e0')

    def update_button_states(self):
        """Update tampilan tombol preview"""
        if self.preview_mode == 'camera':
            self.camera_preview_btn.config(bg='#FF9800', fg='white')
            self.photo_preview_btn.config(bg='#9E9E9E', fg='white')
        else:
            self.camera_preview_btn.config(bg='#9E9E9E', fg='white')
            self.photo_preview_btn.config(bg='#FF9800', fg='white')

    def upload_photo(self):
        """Upload foto untuk analisis"""
        file_path = filedialog.askopenfilename(
            title="Pilih Foto",
            filetypes=[("Image files", "*.jpg *.jpeg *.png *.bmp")]
        )
        if file_path:
            self.uploaded_photo_path = file_path
            self.uploaded_photo_name = Path(file_path).name
            img = cv2.imread(file_path)
            if img is not None:
                self.uploaded_photo = img.copy()
                face = self.detector.get_largest_face(img)
                if face is not None:
                    self.uploaded_face_roi = self.detector.extract_face_roi(self.uploaded_photo, face)
                else:
                    self.uploaded_face_roi = None
                self.show_photo_preview()

    # ========================================================================
    # ANALYSIS
    # ========================================================================

    def start_analysis(self):
        """Memulai analisis"""
        if self.is_analyzing:
            return
        if self.preview_mode == 'camera' and (not self.camera or not self.camera.is_available()):
            messagebox.showerror("Error", "Kamera tidak aktif!")
            return
        if self.preview_mode == 'photo' and self.uploaded_photo is None:
            messagebox.showerror("Error", "Tidak ada foto yang diupload!")
            return

        try:
            distance = float(self.distance_var.get())
        except:
            messagebox.showerror("Error", "Jarak tidak valid!")
            return

        user_type = self.type_var.get()
        expected_name = self.name_var.get()

        self.is_analyzing = True
        self.analyze_btn.config(text="Menganalisis...", state='disabled')
        self.window.update_idletasks()

        self.perform_analysis(distance, user_type, expected_name)

        self.is_analyzing = False
        self.analyze_btn.config(text="▶ Mulai Pengujian", state='normal')

    def perform_analysis(self, distance, user_type, expected_name):
        """Melakukan analisis pengenalan"""
        start_time = time.time()
        face_detected = True
        original_frame = None

        if self.preview_mode == 'camera':
            frame = self.camera.get_frame()
            if frame is None:
                messagebox.showerror("Error", "Gagal mengambil frame!")
                return
            face = self.detector.get_largest_face(frame)
            if face is None:
                face_detected = False
                face_roi = None
            else:
                original_frame = frame.copy()
                face_roi = self.detector.extract_face_roi(frame, face)
        else:
            if self.uploaded_face_roi is None:
                face_detected = False
                face_roi = None
            else:
                face_roi = self.uploaded_face_roi
                original_frame = self.uploaded_photo.copy() if self.uploaded_photo is not None else None

        if not face_detected or face_roi is None:
            recognized_name = "No Face"
            confidence_raw = 100
            confidence_percent = 0.0
            compute_time = (time.time() - start_time) * 1000
        else:
            processed_face = self.detector.preprocess_face(face_roi, equalize_hist=True)

            if processed_face is None:
                recognized_name = "Error"
                confidence_raw = 100
                confidence_percent = 0.0
                compute_time = (time.time() - start_time) * 1000
            else:
                label_id, confidence_raw = self.recognizer.predict_face(processed_face)
                compute_time = (time.time() - start_time) * 1000
                confidence_percent = max(0, min(100, 100 - confidence_raw))

                if confidence_raw < Config.LBPH_THRESHOLD and label_id in self.trainer.labels_dict:
                    recognized_name = self.trainer.labels_dict[label_id]
                else:
                    recognized_name = "Unknown"

        # Simpan foto
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        analysis_img_dir = Config.DATA_DIR / "analysis_images"
        analysis_img_dir.mkdir(parents=True, exist_ok=True)

        original_filename = f"original_{self.analysis_counter:03d}_{timestamp}.jpg"
        original_filepath = analysis_img_dir / original_filename
        if original_frame is not None:
            cv2.imwrite(str(original_filepath), original_frame)
        else:
            original_filepath = ""

        face_filename = f"face_{self.analysis_counter:03d}_{timestamp}.jpg"
        face_filepath = analysis_img_dir / face_filename
        if face_roi is not None:
            cv2.imwrite(str(face_filepath), face_roi)
        else:
            face_filepath = ""

        result = {
            'no': self.analysis_counter,
            'distance': distance,
            'type': user_type,
            'expected_name': expected_name,
            'recognized_name': recognized_name,
            'confidence_raw': round(confidence_raw, 2),
            'confidence_percent': round(confidence_percent, 2),
            'compute_time_ms': round(compute_time, 2),
            'correct': (recognized_name == expected_name) if user_type == "Terdaftar" else (
                    recognized_name in ["Unknown", "No Face"]),
            'face_detected': face_detected,
            'original_image': str(original_filepath),
            'face_image': str(face_filepath)
        }
        self.analysis_data.append(result)
        self.save_analysis_data()

        # Tambahkan ke treeview
        status_code, status_text = self._get_status_label(result)

        expected_display = result['expected_name'][:15] + "..." if len(
            result['expected_name']) > 15 else result['expected_name']
        recognized_display = result['recognized_name'][:15] + "..." if len(
            result['recognized_name']) > 15 else result['recognized_name']

        self.result_tree.insert('', 'end', values=(
            result['no'],
            f"{result['distance']:.1f}m",
            result['type'][:15],
            expected_display,
            recognized_display,
            f"{result['confidence_percent']:.1f}%",
            f"{result['compute_time_ms']:.1f}ms",
            status_text,
            status_code
        ), tags=(status_code.lower(),))

        self.analysis_counter += 1
        self.total_test_label.config(text=f"Total Data: {len(self.analysis_data)} pengujian")
        self.status_label.config(text=f"Status: {status_text} - {recognized_name} ({confidence_percent:.1f}%)",
                                 fg='#4CAF50' if status_code in ['TP', 'TN'] else '#f44336')

    # ========================================================================
    # DATA PERSISTENCE
    # ========================================================================

    def save_analysis_data(self):
        """Menyimpan data analisis ke file JSON"""
        try:
            data_to_save = []
            for result in self.analysis_data:
                data_to_save.append({
                    'no': result['no'],
                    'distance': result['distance'],
                    'type': result['type'],
                    'expected_name': result['expected_name'],
                    'recognized_name': result['recognized_name'],
                    'confidence_raw': result['confidence_raw'],
                    'confidence_percent': result['confidence_percent'],
                    'compute_time_ms': result['compute_time_ms'],
                    'correct': result['correct'],
                    'face_detected': result.get('face_detected', True),
                    'original_image': str(result.get('original_image', '')),
                    'face_image': str(result.get('face_image', ''))
                })

            with open(self.data_file, 'w', encoding='utf-8') as f:
                json.dump({
                    'analysis_counter': self.analysis_counter,
                    'analysis_data': data_to_save
                }, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f"Error saving data: {e}")

    def load_analysis_data(self):
        """Memuat data analisis dari file JSON"""
        if not self.data_file.exists():
            self.analysis_data = []
            self.analysis_counter = 1
            return

        try:
            with open(self.data_file, 'r', encoding='utf-8') as f:
                saved = json.load(f)
                self.analysis_counter = saved.get('analysis_counter', 1)
                self.analysis_data = saved.get('analysis_data', [])
                for result in self.analysis_data:
                    if 'face_detected' not in result:
                        result['face_detected'] = True
        except Exception as e:
            print(f"Error loading data: {e}")
            self.analysis_data = []
            self.analysis_counter = 1

    # ========================================================================
    # DELETE OPERATIONS
    # ========================================================================

    def delete_selected_row(self):
        """Menghapus baris yang dipilih"""
        selected = self.result_tree.selection()
        if selected:
            for item in selected:
                values = self.result_tree.item(item, 'values')
                if values:
                    no = int(values[0])
                    self.analysis_data = [d for d in self.analysis_data if d['no'] != no]
                self.result_tree.delete(item)
            self.total_test_label.config(text=f"Total Data: {len(self.analysis_data)} pengujian")
            self.save_analysis_data()

    def delete_all_rows(self):
        """Menghapus semua baris"""
        if messagebox.askyesno("Konfirmasi", "Hapus semua hasil pengujian?"):
            self.result_tree.delete(*self.result_tree.get_children())
            self.analysis_data = []
            self.analysis_counter = 1
            self.total_test_label.config(text="Total Data: 0 pengujian")
            self.save_analysis_data()

    # ========================================================================
    # EXPORT TO EXCEL
    # ========================================================================

    def export_to_excel(self):
        """Export hasil ke Excel dengan foto dan confusion matrix"""
        if not self.analysis_data:
            messagebox.showwarning("Warning", "Tidak ada data untuk diexport!")
            return

        file_path = filedialog.asksaveasfilename(
            title="Simpan Hasil Analisis",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")]
        )
        if not file_path:
            return

        try:
            from openpyxl.drawing.image import Image as XLImage

            wb = Workbook()
            ws1 = wb.active
            ws1.title = "Hasil Analisis"

            # Header
            headers = ['No', 'Jarak (m)', 'Tipe', 'Nama Diharapkan', 'Nama Dikenali',
                       'Confidence Raw', 'Confidence %', 'Waktu (ms)', 'Status', 'Kode',
                       'Foto Asli', 'Foto Wajah']

            ws1.append(headers)

            header_fill = PatternFill(start_color="FF9800", end_color="FF9800", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            for col in range(1, len(headers) + 1):
                cell = ws1.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center')

            # Isi data
            for i, result in enumerate(self.analysis_data):
                row_num = i + 2
                status_code, status_text = self._get_status_label(result)

                ws1.append([result['no'], result['distance'], result['type'],
                            result['expected_name'], result['recognized_name'],
                            result['confidence_raw'], f"{result['confidence_percent']:.1f}%",
                            result['compute_time_ms'], status_text, status_code,
                            "", ""])

                ws1.row_dimensions[row_num].height = 150

                # Warna baris
                color_map = {
                    "TP": "C8E6C9",
                    "TN": "BBDEFB",
                    "FP": "FFCDD2",
                    "FN": "FFF9C4"
                }
                if status_code in color_map:
                    fill_color = PatternFill(start_color=color_map[status_code],
                                             end_color=color_map[status_code], fill_type="solid")
                    for col in range(1, 11):
                        ws1.cell(row=row_num, column=col).fill = fill_color

                # ===== EMBED FOTO ASLI (PERTAHANKAN ASPECT RATIO) =====
                if result.get('original_image') and Path(result['original_image']).exists():
                    img = XLImage(str(result['original_image']))
                    original = cv2.imread(str(result['original_image']))
                    if original is not None:
                        orig_h, orig_w = original.shape[:2]
                        max_width = 200
                        max_height = 130
                        ratio = min(max_width / orig_w, max_height / orig_h)
                        img.width = int(orig_w * ratio)
                        img.height = int(orig_h * ratio)
                    else:
                        img.width = 200
                        img.height = 112
                    ws1.add_image(img, f'K{row_num}')

                # ===== EMBED FOTO WAJAH (PERTAHANKAN ASPECT RATIO) =====
                if result.get('face_image') and Path(result['face_image']).exists():
                    img = XLImage(str(result['face_image']))
                    face_img = cv2.imread(str(result['face_image']))
                    if face_img is not None:
                        face_h, face_w = face_img.shape[:2]
                        max_face = 120
                        ratio = min(max_face / face_w, max_face / face_h)
                        img.width = int(face_w * ratio)
                        img.height = int(face_h * ratio)
                    else:
                        img.width = 120
                        img.height = 120
                    ws1.add_image(img, f'L{row_num}')

            # Lebar kolom
            col_widths = {'A': 6, 'B': 10, 'C': 14, 'D': 20, 'E': 20,
                          'F': 15, 'G': 15, 'H': 12, 'I': 22, 'J': 8,
                          'K': 28, 'L': 20}
            for col_letter, width in col_widths.items():
                ws1.column_dimensions[col_letter].width = width

            # ===== SHEET 2: CONFUSION MATRIX & METRIK =====
            ws2 = wb.create_sheet("Analisis Statistik")

            ws2['A1'] = "CONFUSION MATRIX"
            ws2['A1'].font = Font(bold=True, size=14, color="333333")
            ws2.merge_cells('A1:D1')

            # Hitung TP, FP, TN, FN
            tp = sum(1 for r in self.analysis_data
                     if r['type'] == 'Terdaftar'
                     and r['recognized_name'] == r['expected_name']
                     and r['recognized_name'] not in ['Unknown', 'No Face', 'Error'])

            fp = sum(1 for r in self.analysis_data
                     if r['type'] == 'Tidak Terdaftar'
                     and r['recognized_name'] not in ['Unknown', 'No Face', 'Error'])

            tn = sum(1 for r in self.analysis_data
                     if r['type'] == 'Tidak Terdaftar'
                     and r['recognized_name'] in ['Unknown', 'No Face', 'Error'])

            fn = sum(1 for r in self.analysis_data
                     if r['type'] == 'Terdaftar'
                     and (r['recognized_name'] != r['expected_name']
                          or r['recognized_name'] in ['Unknown', 'No Face', 'Error']))

            # Confusion Matrix Table
            ws2['A3'] = "Confusion Matrix:"
            ws2['A3'].font = Font(bold=True, size=11)

            cm_headers = ['', 'Predicted Positive\n(Dikenali)', 'Predicted Negative\n(Tidak Dikenali)']
            cm_rows = [
                ['Actual Positive\n(Terdaftar)', f"TP = {tp}", f"FN = {fn}"],
                ['Actual Negative\n(Tidak Terdaftar)', f"FP = {fp}", f"TN = {tn}"]
            ]

            row = 5
            for j, header in enumerate(cm_headers):
                cell = ws2.cell(row=row, column=j + 1, value=header)
                cell.font = Font(bold=True, size=10)
                cell.fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
                cell.alignment = Alignment(horizontal='center', wrap_text=True)

            for i, cm_row in enumerate(cm_rows):
                for j, value in enumerate(cm_row):
                    cell = ws2.cell(row=row + i + 1, column=j + 1, value=value)
                    cell.alignment = Alignment(horizontal='center', wrap_text=True)
                    if i == 0 and j == 1:
                        cell.fill = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")
                    elif i == 0 and j == 2:
                        cell.fill = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")
                    elif i == 1 and j == 1:
                        cell.fill = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")
                    elif i == 1 and j == 2:
                        cell.fill = PatternFill(start_color="BBDEFB", end_color="BBDEFB", fill_type="solid")

            # Keterangan
            row += 4
            ws2.cell(row=row, column=1, value="Keterangan:").font = Font(bold=True, size=11)
            row += 1

            keterangan_data = [
                ("TP (True Positive) - Berhasil Dikenali", tp, "C8E6C9"),
                ("FP (False Positive) - Gagal (Dikenali)", fp, "FFCDD2"),
                ("TN (True Negative) - Berhasil Ditolak", tn, "BBDEFB"),
                ("FN (False Negative) - Gagal (Ditolak)", fn, "FFF9C4"),
            ]

            for label, value, color in keterangan_data:
                cell = ws2.cell(row=row, column=1, value=f"{label}: {value}")
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                row += 1

            # Metrik evaluasi
            row += 1
            ws2.cell(row=row, column=1, value="METRIK EVALUASI").font = Font(bold=True, size=14, color="333333")
            ws2.merge_cells(f'A{row}:D{row}')

            row += 2
            total = tp + tn + fp + fn

            metrics = [
                ("Total Data", total),
                ("Accuracy", f"{(tp + tn) / total * 100:.2f}%" if total > 0 else "N/A"),
                ("Precision", f"{tp / (tp + fp) * 100:.2f}%" if (tp + fp) > 0 else "N/A"),
                ("Recall (Sensitivity)", f"{tp / (tp + fn) * 100:.2f}%" if (tp + fn) > 0 else "N/A"),
                ("Specificity", f"{tn / (tn + fp) * 100:.2f}%" if (tn + fp) > 0 else "N/A"),
                ("F1-Score", f"{2 * tp / (2 * tp + fp + fn) * 100:.2f}%" if (2 * tp + fp + fn) > 0 else "N/A"),
            ]

            for i, (metric_name, metric_value) in enumerate(metrics):
                ws2.cell(row=row + i, column=1, value=metric_name).font = Font(bold=True)
                ws2.cell(row=row + i, column=2, value=metric_value)

            # Lebar kolom Sheet 2
            ws2.column_dimensions['A'].width = 30
            ws2.column_dimensions['B'].width = 22
            ws2.column_dimensions['C'].width = 22
            ws2.column_dimensions['D'].width = 22

            # Tinggi baris untuk confusion matrix
            for r in range(5, 8):
                ws2.row_dimensions[r].height = 35

            wb.save(file_path)
            messagebox.showinfo("Sukses", f"Data berhasil diexport ke:\n{file_path}")
        except Exception as e:
            messagebox.showerror("Error", f"Gagal export data: {str(e)}")

    # ========================================================================
    # CLOSE
    # ========================================================================

    def on_closing(self):
        """Handler saat window ditutup"""
        self.is_running = False

        if self._preview_after_id:
            self.window.after_cancel(self._preview_after_id)
            self._preview_after_id = None

        self.save_analysis_data()
        if self.camera:
            self.camera.stop()
        if self.window:
            self.window.destroy()